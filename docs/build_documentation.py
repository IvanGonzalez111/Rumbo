from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_campaign_assets import generate_campaign_assets


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
BUILD = DOCS / "build"
SOURCE = DOCS / "documentacion-obligatorio.md"
OUTPUT = DOCS / "Rumbo-documentacion-obligatorio.docx"

NAVY = "17223B"
ORANGE = "E86F35"
GREEN = "2F8F68"
SKY = "4A83B8"
ROSE = "B85F68"
CREAM = "F7EFD8"
PAPER = "FFFAF0"
MUTED = "657087"
LINE = "DED0A9"

ARIAL = "/System/Library/Fonts/Supplemental/Arial.ttf"
ARIAL_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(ARIAL_BOLD if bold else ARIAL, size=size)


def cover_crop(image: Image.Image, size: tuple[int, int]) -> Image.Image:
    target_w, target_h = size
    scale = max(target_w / image.width, target_h / image.height)
    resized = image.resize((round(image.width * scale), round(image.height * scale)), Image.Resampling.LANCZOS)
    left = (resized.width - target_w) // 2
    top = (resized.height - target_h) // 2
    return resized.crop((left, top, left + target_w, top + target_h))


def add_gradient_overlay(image: Image.Image, top_alpha: int, bottom_alpha: int) -> Image.Image:
    overlay = Image.new("RGBA", image.size, (10, 25, 50, 0))
    draw = ImageDraw.Draw(overlay)
    height = image.height
    for y in range(height):
        ratio = y / max(1, height - 1)
        alpha = round(top_alpha + (bottom_alpha - top_alpha) * ratio)
        draw.line((0, y, image.width, y), fill=(10, 25, 50, alpha))
    return Image.alpha_composite(image.convert("RGBA"), overlay)


def rounded_label(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, fill: str, text_fill: str) -> None:
    draw.rounded_rectangle(box, radius=18, fill=fill)
    x1, y1, x2, y2 = box
    font = pil_font(24, True)
    bbox = draw.textbbox((0, 0), text, font=font)
    draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, (y1 + y2 - (bbox[3] - bbox[1])) / 2 - 2), text, font=font, fill=text_fill)


def generate_visual_assets() -> None:
    BUILD.mkdir(parents=True, exist_ok=True)
    backdrop = Image.open(ROOT / "public" / "assets" / "rumbo-globe-backdrop.png").convert("RGB")

    # Cover
    cover = add_gradient_overlay(cover_crop(backdrop, (1500, 2000)), 80, 35)
    draw = ImageDraw.Draw(cover)
    draw.rounded_rectangle((105, 105, 285, 285), radius=90, fill=f"#{ORANGE}", outline=f"#{PAPER}", width=6)
    draw.text((155, 137), "R", font=pil_font(100, True), fill=f"#{PAPER}")
    draw.text((105, 350), "RUMBO", font=pil_font(150, True), fill=f"#{PAPER}")
    draw.rounded_rectangle((110, 525, 570, 537), radius=6, fill=f"#{ORANGE}")
    draw.text((105, 590), "Memoria de diseño, desarrollo", font=pil_font(46, True), fill=f"#{PAPER}")
    draw.text((105, 652), "e integración de inteligencia artificial", font=pil_font(42), fill=f"#{PAPER}")
    draw.text((105, 765), "OBLIGATORIO 2 · DISEÑO INTERACTIVO", font=pil_font(28, True), fill="#F0B48E")
    draw.text((105, 825), "Iván Gonzales · Julio 2026", font=pil_font(28), fill="#E8EDF5")
    draw.text((105, 1830), "Convertí tu viaje en una aventura interactiva.", font=pil_font(34, True), fill=f"#{PAPER}")
    cover.convert("RGB").save(BUILD / "cover.png", quality=94)

    # Architecture diagram
    diagram = Image.new("RGB", (1800, 650), f"#{PAPER}")
    draw = ImageDraw.Draw(diagram)
    draw.text((80, 50), "Cómo viajan los datos y la IA", font=pil_font(50, True), fill=f"#{NAVY}")
    nodes = [
        ("USUARIO", "acciones y contexto", SKY),
        ("FRONTEND", "HTML · CSS · JS", ORANGE),
        ("BACKEND", "Node · Express", NAVY),
        ("PROXY IA", "Gemini Vertex", GREEN),
        ("VALIDACIÓN", "JSON normalizado", ROSE),
        ("PERSISTENCIA", "Supabase\nPostgres + Storage", SKY),
    ]
    x_positions = [75, 365, 655, 945, 1235, 1525]
    for idx, ((title, subtitle, color), x) in enumerate(zip(nodes, x_positions)):
        box = (x, 250, x + 220, 430)
        draw.rounded_rectangle(box, radius=22, fill="#FFFFFF", outline=f"#{color}", width=5)
        draw.rounded_rectangle((x, 250, x + 220, 298), radius=18, fill=f"#{color}")
        draw.rectangle((x, 280, x + 220, 300), fill=f"#{color}")
        title_bbox = draw.textbbox((0, 0), title, font=pil_font(23, True))
        draw.text((x + 110 - (title_bbox[2] - title_bbox[0]) / 2, 262), title, font=pil_font(23, True), fill="#FFFFFF")
        wrapped = subtitle.replace(" · ", "\n")
        draw.multiline_text((x + 110, 337), wrapped, font=pil_font(23), fill=f"#{NAVY}", anchor="mm", align="center", spacing=8)
        if idx < len(nodes) - 1:
            x1 = x + 230
            x2 = x_positions[idx + 1] - 10
            y = 340
            draw.line((x1, y, x2, y), fill=f"#{MUTED}", width=5)
            draw.polygon([(x2, y), (x2 - 18, y - 12), (x2 - 18, y + 12)], fill=f"#{MUTED}")
    draw.text((80, 535), "La credencial permanece en el servidor. El navegador nunca llama directamente al modelo.", font=pil_font(28), fill=f"#{MUTED}")
    diagram.save(ASSETS / "07-arquitectura-ia.png", quality=95)

    # Communication pieces
    feed = add_gradient_overlay(cover_crop(backdrop, (1080, 1080)), 105, 80)
    draw = ImageDraw.Draw(feed)
    rounded_label(draw, (70, 68, 280, 126), "RUMBO", f"#{ORANGE}", f"#{PAPER}")
    draw.text((70, 210), "NO HAGAS", font=pil_font(82, True), fill=f"#{PAPER}")
    draw.text((70, 300), "SOLO FOTOS.", font=pil_font(82, True), fill=f"#{PAPER}")
    draw.text((70, 420), "DESBLOQUEÁ", font=pil_font(77, True), fill="#F6A06F")
    draw.text((70, 505), "RECUERDOS.", font=pil_font(77, True), fill="#F6A06F")
    draw.rounded_rectangle((70, 870, 780, 980), radius=24, fill=f"#{PAPER}")
    draw.text((105, 900), "Misiones · Sellos · Bitácora", font=pil_font(37, True), fill=f"#{NAVY}")
    feed.convert("RGB").save(ASSETS / "08-pieza-feed.png", quality=94)

    story = add_gradient_overlay(cover_crop(backdrop, (1080, 1920)), 110, 75)
    draw = ImageDraw.Draw(story)
    rounded_label(draw, (70, 70, 280, 130), "RUMBO", f"#{ORANGE}", f"#{PAPER}")
    draw.text((70, 260), "TU VIAJE,", font=pil_font(91, True), fill=f"#{PAPER}")
    draw.text((70, 360), "UNA MISIÓN", font=pil_font(91, True), fill=f"#{PAPER}")
    draw.text((70, 460), "A LA VEZ.", font=pil_font(91, True), fill="#F6A06F")
    draw.text((70, 1580), "Explorá. Guardá. Recordá.", font=pil_font(43, True), fill=f"#{PAPER}")
    draw.rounded_rectangle((70, 1660, 840, 1790), radius=28, fill=f"#{ORANGE}")
    draw.text((112, 1698), "CONVERTÍ TU VIAJE EN AVENTURA", font=pil_font(30, True), fill=f"#{PAPER}")
    story.convert("RGB").save(ASSETS / "09-pieza-story.png", quality=94)
    generate_campaign_assets()


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, name="Arial", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), SKY)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|https?://[^\s)]+)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Courier New", 9.5, NAVY)
            run.font.highlight_color = None
        else:
            add_hyperlink(paragraph, token, token)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])
    for run in paragraph.runs:
        if run.font.name is None:
            set_run_font(run, "Arial", 10.5, NAVY)


def style_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.22

    tokens = {
        "Heading 1": (18, ORANGE, 18, 8),
        "Heading 2": (14, NAVY, 15, 6),
        "Heading 3": (11.5, SKY, 11, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False


def set_header_footer(section) -> None:
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    left = p.add_run("RUMBO")
    set_run_font(left, size=8.5, color=ORANGE, bold=True)
    right = p.add_run("  ·  MEMORIA DEL PROYECTO")
    set_run_font(right, size=8.5, color=MUTED, bold=True)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), LINE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Obligatorio 2 · Diseño Interactivo · Iván Gonzales")
    for run in fp.runs:
        set_run_font(run, size=8, color=MUTED)


def add_picture(doc: Document, filename: str, caption: str, width=6.35) -> None:
    path = ASSETS / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(caption)


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.05
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EFF2F5")
    p_pr.append(shd)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, "Courier New", 8.3, NAVY)
        if index < len(lines) - 1:
            run.add_break()


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    header = rows[0]
    body = rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next = True
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = text.strip()
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8.5, color=PAPER, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(body):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value.strip()
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx])
            if row_index % 2:
                set_cell_shading(cells[idx], "F5F0E4")
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    set_run_font(run, size=8.2, color=NAVY)
    for cell in table.rows[0].cells:
        set_cell_margins(cell)
    widths = [1600, 2200, 2800, 2760] if len(header) == 4 else [round(9360 / len(header))] * len(header)
    widths[-1] += 9360 - sum(widths)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_section_visuals(doc: Document, heading: str) -> None:
    visuals = {
        "8. Bocetos, wireframes y prototipo": [
            ("01-dashboard.png", "Figura 1. Dashboard con portada narrativa, estado y progreso del viaje."),
        ],
        "9. Arquitectura de información y flujo": [
            ("02-misiones.png", "Figura 2. Próxima acción, recorrido y misiones en una misma vista orientada a la tarea."),
        ],
        "10. Definición de tecnologías": [
            ("06-crear-viaje.png", "Figura 3. Formulario personalizado con grupo, estilos, energía, fechas y destino."),
        ],
        "13. Integración de inteligencia artificial": [
            ("07-arquitectura-ia.png", "Figura 4. Arquitectura de la integración con el proxy de IA."),
        ],
        "17. Estrategia de comunicación": [
            ("08-pieza-feed.png", "Figura 5. Pieza para feed: mensaje principal de lanzamiento."),
            ("09-pieza-story.png", "Figura 6. Pieza vertical para stories y reels."),
            ("10-carrusel-flujo.png", "Figura 7. Carrusel explicativo del recorrido completo de Rumbo."),
            ("13-mockup-redes.png", "Figura 8. Adaptación de la campaña a feed, story y carrusel."),
            ("14-mockup-via-publica.png", "Figura 9. Aplicación del concepto en un soporte de vía pública."),
            ("15-tablero-campana.png", "Figura 10. Sistema visual y familia completa de piezas de lanzamiento."),
        ],
        "18. Pruebas y validación": [
            ("03-completar-mision.png", "Figura 11. Flujo de tres pasos para guardar evidencia, emoción e historia."),
        ],
        "20. Estado actual y próximos pasos": [
            ("04-pasaporte.png", "Figura 12. Pasaporte como colección visual de sellos."),
            ("05-bitacora.png", "Figura 13. Bitácora final como resultado narrativo del viaje."),
        ],
    }
    for filename, caption in visuals.get(heading, []):
        if filename == "09-pieza-story.png":
            width = 3.5
        elif filename in {"10-carrusel-flujo.png", "13-mockup-redes.png", "14-mockup-via-publica.png", "15-tablero-campana.png"}:
            width = 6.35
        elif filename == "05-bitacora.png":
            width = 3.2
        elif filename == "01-dashboard.png":
            width = 4.45
        else:
            width = 6.35
        add_picture(doc, filename, caption, width=width)


def parse_markdown(doc: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1. "))
    lines = lines[start:]
    paragraph_buffer: list[str] = []
    code_buffer: list[str] = []
    in_code = False
    table_buffer: list[list[str]] = []
    major_breaks = set()

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline(p, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    for raw in lines + [""]:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_paragraph()
            flush_table()
            if in_code:
                add_code_block(doc, code_buffer)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buffer.append(line)
            continue
        if line.startswith("|") and line.endswith("|"):
            flush_paragraph()
            table_buffer.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        flush_table()
        if not line:
            flush_paragraph()
            continue
        if line == "---":
            flush_paragraph()
            continue
        if line.startswith("## "):
            flush_paragraph()
            text = line[3:]
            if any(text.startswith(prefix) for prefix in major_breaks) and len(doc.paragraphs) > 3:
                doc.add_page_break()
            p = doc.add_paragraph(text, style="Heading 1")
            p.paragraph_format.keep_with_next = True
            add_section_visuals(doc, text)
            continue
        if line.startswith("### "):
            flush_paragraph()
            p = doc.add_paragraph(line[4:], style="Heading 2")
            p.paragraph_format.keep_with_next = True
            continue
        if line.startswith("#### "):
            flush_paragraph()
            p = doc.add_paragraph(line[5:], style="Heading 3")
            p.paragraph_format.keep_with_next = True
            continue
        if line.startswith("> "):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(9)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "FBE7D8")
            p_pr.append(shd)
            add_inline(p, line[2:])
            continue
        if re.match(r"^- ", line):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            continue
        numbered = re.match(r"^(\d+)\. (.+)", line)
        if numbered:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(4)
            number_run = p.add_run(f"{numbered.group(1)}.  ")
            set_run_font(number_run, size=10.5, color=ORANGE, bold=True)
            add_inline(p, numbered.group(2))
            continue
        paragraph_buffer.append(line)


def add_contents(doc: Document) -> None:
    p = doc.add_paragraph("Contenido", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    groups = [
        "01 · Fundamentos: modalidad, problema, público y persona",
        "02 · Investigación, ideación y evolución del prototipo",
        "03 · Tecnología, diseño y desarrollo",
        "04 · Integración de IA, prompts y seguridad",
        "05 · Automatizaciones, comunicación y piezas",
        "06 · Pruebas, aprendizajes y próximos pasos",
        "07 · Anexos y ejecución",
    ]
    for item in groups:
        row = doc.add_paragraph()
        row.paragraph_format.space_after = Pt(7)
        number, text = item.split(" · ", 1)
        num_run = row.add_run(number + "  ")
        set_run_font(num_run, size=9.5, color=ORANGE, bold=True)
        text_run = row.add_run(text)
        set_run_font(text_run, size=10.5, color=NAVY)
    doc.add_page_break()


def build_document() -> None:
    generate_visual_assets()
    doc = Document()
    style_document(doc)

    cover_section = doc.sections[0]
    cover_section.top_margin = Inches(0.5)
    cover_section.bottom_margin = Inches(0.5)
    cover_section.left_margin = Inches(0.5)
    cover_section.right_margin = Inches(0.5)
    cover_section.header_distance = Inches(0)
    cover_section.footer_distance = Inches(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(BUILD / "cover.png"), width=Inches(7.5), height=Inches(10))

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.header.is_linked_to_previous = False
    body_section.footer.is_linked_to_previous = False
    set_header_footer(body_section)

    add_contents(doc)
    parse_markdown(doc)

    core = doc.core_properties
    core.title = "Rumbo - Memoria de diseño, desarrollo e integración de IA"
    core.subject = "Obligatorio 2 de Diseño Interactivo"
    core.author = "Iván Gonzales"
    core.keywords = "Rumbo, Diseño Interactivo, IA, UX, viajes, Gemini"
    core.comments = "MVP funcional documentado el 20 de julio de 2026."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
